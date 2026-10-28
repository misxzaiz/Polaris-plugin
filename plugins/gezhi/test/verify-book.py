# -*- coding: utf-8 -*-
"""
格致·产物级验收：用 openpyxl 打开引擎产出的 .xlsx，断言它是一份
真正的 Excel（样式/公式/冻结/筛选/条件格式/类型/注入防护全部成立）。
用法: python test/verify-book.py <xlsx路径>
"""
import sys
from openpyxl import load_workbook

path = sys.argv[1]
wb = load_workbook(path)

fails = []


def check(cond, msg):
    if not cond:
        fails.append(msg)


# ── 结构 ────────────────────────────────────────────────────────────────────
check(wb.sheetnames == ['采购明细', '汇总'], f'sheet 名应为 [采购明细, 汇总]，实为 {wb.sheetnames}')
ws = wb['采购明细']

# ── 标题行（合并 A1:F1）────────────────────────────────────────────────────
check('A1:F1' in [str(r) for r in ws.merged_cells.ranges], 'caption 未合并 A1:F1')
check(ws['A1'].value == '生成于 2026-10-29', f'caption 内容错：{ws["A1"].value}')

# ── 表头（行 2）：值 + 白字 + 主题填充 ─────────────────────────────────────
headers = [ws.cell(row=2, column=c).value for c in range(1, 7)]
check(headers == ['品名', '数量', '单价', '小计', '到货日', '毛利率'], f'表头错：{headers}')
h1 = ws['A2']
check(h1.font.bold, '表头未加粗')
check(h1.font.color.rgb in ('FFFFFFFF', '00FFFFFF'), f'表头字色应白，实为 {h1.font.color.rgb}')
check(h1.fill.fgColor.rgb == 'FF4338CA', f'indigo 表头填充错：{h1.fill.fgColor.rgb}')

# ── 数据与类型 ─────────────────────────────────────────────────────────────
check(ws['B3'].value == 12 and isinstance(ws['B3'].value, (int, float)), f'数量应为数字：{ws["B3"].value!r}')
check(ws['C3'].value == 3.5, f'单价应为数字：{ws["C3"].value!r}')
check(ws['C3'].number_format == '#,##0.00', f'money 格式错：{ws["C3"].number_format}')
import datetime
check(isinstance(ws['E3'].value, datetime.datetime), f'日期应为 datetime：{ws["E3"].value!r}')
check(ws['E3'].number_format == 'yyyy-mm-dd', f'date 格式错：{ws["E3"].number_format}')
# 百分比：'15%' 字符串 → 0.15（0.0% 格式）
check(abs((ws['F3'].value or 0) - 0.15) < 1e-9, f'percent 应转 0.15：{ws["F3"].value!r}')
check(ws['F3'].number_format == '0.0%', f'percent 格式错：{ws["F3"].number_format}')
# GB18030/乱码无关，但中文必须原样
check(ws['A4'].value == '香蕉', f'中文值错：{ws["A4"].value!r}')

# ── 行公式（小计列 D）─────────────────────────────────────────────────────
check(ws['D3'].data_type == 'f', f'小计应为公式，data_type={ws["D3"].data_type}')
check(str(ws['D3'].value).replace('$', '').upper() == '=B3*C3', f'公式错：{ws["D3"].value}')

# ── 公式注入防护（=危险公式 必须是文本，不是公式）────────────────────────
a5 = ws['A5']
check(a5.data_type == 's', f'以=开头的文本必须按文本存储，data_type={a5.data_type}')
check(a5.value == '=危险公式', f'注入防护后显示值应原样：{a5.value!r}')

# ── 汇总行（真实 =SUM 公式 + 双线上边框）──────────────────────────────────
check(ws['B7'].data_type == 'f', f'汇总应为公式：{ws["B7"].data_type}')
check('SUM(B3:B6)' in str(ws['B7'].value).replace('$', '').upper(), f'SUM 公式错：{ws["B7"].value}')
check(ws['A7'].value == '合计', f'汇总标签错：{ws["A7"].value}')
check(ws['B7'].border.top.style == 'double', f'汇总行应为双线上边框：{ws["B7"].border.top.style}')
check(ws['B7'].font.bold, '汇总行未加粗')

# ── 冻结 / 筛选 ────────────────────────────────────────────────────────────
check(ws.freeze_panes == 'A3', f'冻结应在 A3（caption+表头之下），实为 {ws.freeze_panes}')
check(ws.auto_filter.ref == 'A3:F6', f'筛选范围错：{ws.auto_filter.ref}')

# ── 条件格式（2 条 expression 规则）───────────────────────────────────────
cfs = list(ws.conditional_formatting)
total_rules = sum(len(cf.rules) for cf in cfs)
check(total_rules == 2, f'条件格式应有 2 条，实为 {total_rules}')
found_gt = False
found_contains = False
for cf in cfs:
    for rule in cf.rules:
        f0 = (rule.formula[0] if getattr(rule, 'formula', None) else '').replace('$', '').replace(' ', '')
        if 'D3>50' in f0:
            found_gt = True
            check('FEF3C7' in (rule.dxf.fill.bgColor.rgb or ''), f'warn 填充色错：{rule.dxf.fill.bgColor.rgb}')
        if 'ISNUMBER(SEARCH' in f0 and '即期' in f0:
            found_contains = True
check(found_gt, '缺 D3>50 条件规则')
check(found_contains, '缺 contains(即期) 条件规则')

# ── 列宽（应按内容计算而非默认 13）────────────────────────────────────────
w = ws.column_dimensions['A'].width
check(w and w > 8, f'A 列宽未按内容计算：{w}')

# ── 第二个 sheet（无 caption：表头在行 1，数据行 2 起）─────────────────────
ws2 = wb['汇总']
check(ws2['A1'].value == '指标', f'汇总 sheet 表头错：{ws2["A1"].value}')
check(ws2['A2'].value == '品项数', f'汇总 sheet 数据错：{ws2["A2"].value}')
check(ws2.freeze_panes == 'A2', f'汇总 sheet 冻结应在 A2：{ws2.freeze_panes}')
check(ws2.auto_filter.ref == 'A2:B3', f'汇总 sheet 筛选错：{ws2.auto_filter.ref}')

if fails:
    print('FAIL')
    for f in fails:
        print(' -', f)
    sys.exit(1)
print('PASS: 产物级验收全部成立（结构/样式/公式/冻结/筛选/条件格式/注入防护/类型）')
